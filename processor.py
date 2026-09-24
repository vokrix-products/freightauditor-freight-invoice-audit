import csv
import io
import re
from datetime import date, datetime
from typing import Any, Dict, List

import pdfplumber
from openpyxl import load_workbook

from llm_extractor import extract_from_unstructured_text

STATUS_MISSING = "missing:critical"
STATUS_EXPIRED = "expired:warning"
STATUS_VALID = "valid:good"
STATUS_FLAGGED = "flagged:critical"
STATUS_CONTRACT_REVIEW = "contract-review:warning"
STATUS_UNMAPPED = "unmapped:warning"

FLOAT_TOLERANCE = 0.01

NUMERIC_FIELDS = {
    "actual_weight",
    "billable_weight",
    "freight_charge",
    "fuel_surcharge",
    "accessorial_charges",
    "total_charges",
    "minimum_charge",
    "base_rate",
    "expected_freight_charge",
    "expected_fuel_surcharge",
    "expected_accessorial_charges",
    "total_expected_charge",
    "variance",
    "overcharge_amount",
}

KEY_ALIASES_RAW = {
    "invoice_number": ["invoice number", "invoice no", "invoice #", "invoice_num", "invoiceid", "inv no", "invoice_number"],
    "carrier_name": ["carrier name", "carrier", "supplier", "vendor name", "vendor", "carrier_name"],
    "ship_date": ["ship date", "shipment date", "shipping date", "date shipped", "shipdate", "ship_date"],
    "bill_of_lading_pro_number": ["bill of lading", "pro number", "pro no", "bol", "bol/pro", "bol_number", "pro_number", "bill of lading / pro number"],
    "origin_location": ["origin location", "origin", "origin city", "origin address", "origin city/state/postal code", "origin_zip", "origin zone"],
    "destination_location": ["destination location", "destination", "destination city", "destination address", "destination city/state/postal code", "destination_zip", "destination zone"],
    "actual_weight": ["actual weight", "actual weight (lbs)", "actual_weight", "weight"],
    "billable_weight": ["billable weight", "billed weight", "billable_weight", "billing weight"],
    "freight_charge": ["freight charge", "freight", "freight charges", "freight_charge", "linehaul", "line haul", "linehaul charge", "line haul charge"],
    "fuel_surcharge": ["fuel surcharge", "fuel", "fuel_surcharge", "fsc"],
    "accessorial_codes": ["accessorial code", "accessorial codes", "accessorial_code"],
    "accessorial_descriptions": ["accessorial description", "accessorial descriptions", "accessorial_desc"],
    "accessorial_charges": ["accessorial charge", "accessorial charges", "accessorial_charge"],
    "total_charges": ["total charges", "total charge", "total", "invoice total", "total_charges", "price", "amount", "total_amount", "total amount", "grand total", "invoice amount"],
    "payment_due_date": ["payment due date", "due date", "payment_due_date", "due_date", "invoice due date"],
    "currency": ["currency", "curr"],
    "invoice_line_items": ["invoice line item details", "line item details", "line_items", "invoice line-item details", "invoice_line_items"],
    "source_file_name": ["source file name", "source filename", "filename"],
    "contract_rate_sheet_identifier": ["contract/rate-sheet identifier", "contract id", "rate sheet id", "rate_sheet_id", "contract_rate_sheet_identifier", "contract_reference"],
    "effective_date": ["effective date", "effective", "eff date", "effective_date", "rate_sheet_effective_date", "rate sheet effective date"],
    "expiration_date": ["expiration date", "expiration", "exp date", "expiration_date", "rate_sheet_expiration_date", "rate sheet expiration date"],
    "origin_zone_zip_postal": ["origin zone/zip/postal code", "origin zone", "origin_zip", "origin postal", "origin_zone", "origin_zone_zip_postal"],
    "destination_zone_zip_postal": ["destination zone/zip/postal code", "destination zone", "destination_zip", "destination postal", "destination_zone", "destination_zone_zip_postal"],
    "freight_class_commodity": ["freight class/commodity", "freight class", "commodity", "class", "freight_class", "product", "freight_class_commodity"],
    "rate_basis": ["rate basis", "rate_basis", "basis"],
    "minimum_charge": ["minimum charge", "minimum_charge", "min charge"],
    "base_rate": ["base rate", "base_rate", "rate"],
    "fuel_surcharge_table": ["fuel surcharge table/percentage schedule", "fuel surcharge table", "fuel schedule", "fuel percentage", "fuel_surcharge_table", "fuel_surcharge_schedule", "fuel surcharge schedule"],
    "accessorial_rule_code_description_amount": ["accessorial rule/code/description/amount", "accessorial rule", "accessorial code"],
    "expected_freight_charge": ["expected freight charge", "expected freight", "expected_freight_charge"],
    "expected_fuel_surcharge": ["expected fuel surcharge", "expected fuel", "expected_fuel_surcharge"],
    "expected_accessorial_charges": ["expected accessorial charges", "expected accessorial", "expected_accessorial_charges"],
    "total_expected_charge": ["total expected charge", "total expected", "total_expected_charge"],
    "variance": ["variance"],
    "overcharge_amount": ["overcharge amount", "overcharge"],
    "matched_rate_line_reference": ["matched rate line reference", "matched rate line"],
    "invoice_status": ["invoice status"],
    "rate_sheet_status": ["rate-sheet status", "rate sheet status"],
    "mapping_template_version": ["mapping template version"],
}


# Fields that prove a document really is an invoice / rate sheet. Detection scores
# these by VALUE. It must not use key presence: the extractor returns the whole
# schema with explicit nulls, so every document has an invoice_number key.
INVOICE_MARKER_FIELDS = (
    "invoice_number",
    "invoice_date",
    "ship_date",
    "freight_charge",
    "invoice_line_items",
    "total_charges",
    "bill_of_lading_pro_number",
)

RATE_SHEET_MARKER_FIELDS = (
    "lanes",
    "contract_rate_sheet_identifier",
    "effective_date",
    "expiration_date",
    "rate_basis",
    "base_rate",
    "minimum_charge",
    "fuel_surcharge_table",
)


def _token(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


KEY_ALIASES = {
    canonical: {_token(alias) for alias in aliases}
    for canonical, aliases in KEY_ALIASES_RAW.items()
}


def _canonical_key(raw_key: str) -> str:
    token = _token(str(raw_key))
    for canonical, aliases in KEY_ALIASES.items():
        if token in aliases:
            return canonical
    return str(raw_key).strip().lower().replace(" ", "_")


def _clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in {"none", "null", "nan", ""}:
            return None
    return value


def _to_float(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("$", "").replace(",", "").strip()
    if not text or text.lower() in {"none", "null", "nan", ""}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _has_value(record: Dict[str, Any], field: str) -> bool:
    """True when a field carries real content. Nulls, blanks and empty
    containers all mean the document did not actually provide that field."""
    value = record.get(field)
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    if isinstance(value, (list, dict)):
        return len(value) > 0
    if isinstance(value, bool):
        return value
    return True


def _parse_date(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.lower() in {"none", "null", "nan", ""}:
        return None
    for fmt in (
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m-%d-%Y",
        "%d-%b-%Y",
        "%b %d, %Y",
        "%B %d, %Y",
        "%Y/%m/%d",
        "%Y-%m-%d %I:%M %p",
        "%b %d, %Y %I:%M %p",
        "%B %d, %Y %I:%M %p",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except Exception:
        return None


def _iso_date(value: Any) -> Any:
    parsed = _parse_date(value)
    if parsed:
        return parsed.isoformat()
    return None


def _heuristic_carrier(text: str) -> Any:
    match = re.search(r"(?:carrier|vendor|supplier|carrier name|vendor name)\s*[:|-]\s*([A-Za-z0-9 .&]+)", text, re.I)
    if match:
        return match.group(1).strip()
    return None


def _normalize_row(row: Dict[str, Any]) -> Dict[str, Any]:
    if "raw_text" in row:
        raw = row.get("raw_text") or ""
        return {
            "_type": "unknown",
            "raw_text": raw,
            "carrier_name": _heuristic_carrier(raw),
        }

    normalized: Dict[str, Any] = {}
    for raw_key, value in row.items():
        canonical = _canonical_key(raw_key)
        normalized[canonical] = _clean_value(value)

    for field in NUMERIC_FIELDS:
        if field in normalized:
            normalized[field] = _to_float(normalized[field])

    return normalized


def _expand_rate_lanes(record: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Split a rate sheet with a lanes[] array into one record per rate line.

    A rate sheet is audited line by line, and the per-row duplicate check in
    _assign_status already assumes one row per lane. Every line inherits the
    sheet-level fields (carrier, effective/expiration date, contract id); lane
    fields win where they carry a value.
    """
    lanes = record.get("lanes")
    if not isinstance(lanes, list) or not lanes:
        return [record]

    sheet_fields = {key: value for key, value in record.items() if key != "lanes"}
    expanded: List[Dict[str, Any]] = []

    for index, lane in enumerate(lanes):
        if not isinstance(lane, dict):
            continue
        merged = dict(sheet_fields)
        for raw_key, value in lane.items():
            canonical = _canonical_key(raw_key)
            cleaned = _clean_value(value)
            if cleaned is not None:
                merged[canonical] = cleaned
        for field in NUMERIC_FIELDS:
            if field in merged:
                merged[field] = _to_float(merged[field])
        merged["_rate_line_index"] = index
        merged["_document_type_hint"] = "rate_sheet"
        expanded.append(merged)

    return expanded or [record]


def _detect_document_type(record: Dict[str, Any]) -> str:
    invoice_score = sum(1 for field in INVOICE_MARKER_FIELDS if _has_value(record, field))
    rate_score = sum(1 for field in RATE_SHEET_MARKER_FIELDS if _has_value(record, field))

    if rate_score > invoice_score:
        return "rate_sheet"
    if invoice_score:
        return "invoice"
    if rate_score:
        return "rate_sheet"
    if _has_value(record, "carrier_name"):
        return "invoice"
    return "unknown"


def _assign_status(record: Dict[str, Any], all_rows: List[Dict[str, Any]]):
    doc_type = record.get("_type") or "unknown"

    if doc_type == "invoice":
        required = ["invoice_number", "carrier_name", "ship_date", "freight_charge"]
        missing = [field for field in required if record.get(field) in (None, "")]
        if missing:
            return STATUS_MISSING, [f"missing required fields: {', '.join(missing)}"]

        invoice_no = str(record.get("invoice_number") or "").strip()
        duplicate_count = sum(
            1
            for row in all_rows
            if str(row.get("invoice_number") or "").strip() == invoice_no
        )
        if invoice_no and duplicate_count > 1:
            return STATUS_FLAGGED, ["duplicate invoice number in file"]

        total = record.get("total_charges")
        component_fields = ["freight_charge", "fuel_surcharge", "accessorial_charges"]
        components = [record.get(field) for field in component_fields if record.get(field) is not None]
        if total is not None and components:
            expected_total = sum(components)
            if abs(float(total) - expected_total) > FLOAT_TOLERANCE:
                return STATUS_FLAGGED, ["total charges do not match freight+fuel+accessorial"]

        variance_pairs = [
            ("expected_freight_charge", "freight_charge"),
            ("expected_fuel_surcharge", "fuel_surcharge"),
            ("expected_accessorial_charges", "accessorial_charges"),
            ("total_expected_charge", "total_charges"),
        ]
        for expected_field, actual_field in variance_pairs:
            expected = record.get(expected_field)
            actual = record.get(actual_field)
            if expected is not None and actual is not None:
                difference = float(actual) - float(expected)
                if abs(difference) > FLOAT_TOLERANCE:
                    return STATUS_FLAGGED, [f"variance for {expected_field}: {difference:.2f}"]

        return STATUS_VALID, []

    if doc_type == "rate_sheet":
        required = ["carrier_name", "effective_date", "expiration_date", "rate_basis", "base_rate"]
        missing = [field for field in required if record.get(field) in (None, "")]
        if missing:
            return STATUS_MISSING, [f"missing required rate sheet fields: {', '.join(missing)}"]

        expiration = _parse_date(record.get("expiration_date"))
        if expiration and expiration < date.today():
            return STATUS_EXPIRED, ["rate sheet expired"]

        base_rate = record.get("base_rate")
        minimum_charge = record.get("minimum_charge")
        if base_rate is not None and minimum_charge is not None:
            if float(minimum_charge) > float(base_rate) * 3:
                return STATUS_CONTRACT_REVIEW, ["minimum charge is above base-rate threshold"]

        rate_key = (
            record.get("origin_zone_zip_postal"),
            record.get("destination_zone_zip_postal"),
            record.get("freight_class_commodity"),
        )
        if rate_key[0] is not None or rate_key[1] is not None:
            duplicate_count = sum(
                1
                for row in all_rows
                if (
                    row.get("origin_zone_zip_postal"),
                    row.get("destination_zone_zip_postal"),
                    row.get("freight_class_commodity"),
                )
                == rate_key
            )
            if duplicate_count > 1:
                return STATUS_CONTRACT_REVIEW, ["duplicate rate line for same origin/destination/class"]

        rate_basis = str(record.get("rate_basis") or "").lower()
        if rate_basis not in {"flat", "minimum charge", "per unit"} and not record.get("fuel_surcharge_table"):
            return STATUS_CONTRACT_REVIEW, ["missing fuel surcharge schedule"]

        return STATUS_VALID, []

    return STATUS_UNMAPPED, ["unable to identify document type"]


def _try_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        if text.strip():
            return [{"raw_text": text}]
    except Exception:
        pass
    return []


def _try_excel(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        return []

    rows: List[Dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        raw_rows = list(worksheet.iter_rows(values_only=True))
        if not raw_rows:
            continue

        headers: List[str] = []
        for header in raw_rows[0]:
            if header is None or str(header).strip() == "":
                headers.append("")
            else:
                headers.append(str(header).strip().lower())

        if not any(headers):
            continue

        for values in raw_rows[1:]:
            if not any(value is not None and str(value).strip() != "" for value in values):
                continue
            row: Dict[str, Any] = {}
            for idx, value in enumerate(values):
                key = headers[idx] if idx < len(headers) and headers[idx] else f"column_{idx}"
                row[key] = value
            rows.append(row)

    return rows


def _try_text_or_csv(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        text = file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return []

    if not text.strip():
        return []

    sample = text[:512]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = None

    if dialect:
        try:
            reader = csv.DictReader(io.StringIO(text), dialect=dialect)
            rows: List[Dict[str, Any]] = []
            for row in reader:
                if any(value is not None and str(value).strip() != "" for value in row.values()):
                    cleaned: Dict[str, Any] = {}
                    for key, value in row.items():
                        if key is None:
                            continue
                        cleaned[str(key).strip().lower()] = value
                    rows.append(cleaned)
            if rows:
                return rows
        except Exception:
            pass

    return [{"raw_text": text}]


def _extract_rows(file_bytes: bytes) -> List[Dict[str, Any]]:
    rows = _try_pdf(file_bytes)
    if rows:
        return rows

    rows = _try_excel(file_bytes)
    if rows:
        return rows

    rows = _try_text_or_csv(file_bytes)
    if rows:
        return rows

    return []


def process_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    raw_rows = _extract_rows(file_bytes)

    normalized_rows: List[Dict[str, Any]] = []
    for row in raw_rows:
        if "raw_text" in row:
            extracted = extract_from_unstructured_text(row["raw_text"])
            if extracted:
                for item in extracted:
                    normalized = _normalize_row(item)
                    if normalized:
                        normalized_rows.append(normalized)
            else:
                normalized = _normalize_row(row)
                if normalized:
                    normalized_rows.append(normalized)
        else:
            normalized = _normalize_row(row)
            if normalized:
                normalized_rows.append(normalized)

    normalized_rows = [row for row in normalized_rows if row]

    expanded_rows: List[Dict[str, Any]] = []
    for row in normalized_rows:
        expanded_rows.extend(_expand_rate_lanes(row))
    normalized_rows = expanded_rows

    output: List[Dict[str, Any]] = []
    for record in normalized_rows:
        forced_type = record.pop("_document_type_hint", None)
        document_type = forced_type or _detect_document_type(record)
        record["_type"] = document_type

        status, notes = _assign_status(record, normalized_rows)

        title = (
            record.get("carrier_name")
            or record.get("supplier")
            or record.get("vendor_name")
            or "Unknown Vendor"
        )
        due_date_value = record.get("payment_due_date") or record.get("due_date")
        if due_date_value is None and document_type == "rate_sheet":
            # The "Due / Expires" column and the Upcoming Expirations widget both
            # read due_date. A rate sheet has no payment due date, so without this
            # the one document type that actually expires could never surface.
            due_date_value = record.get("expiration_date")
        due_date = _iso_date(due_date_value)

        details: Dict[str, Any] = {}
        for key, value in record.items():
            if key in {
                "carrier_name",
                "supplier",
                "vendor_name",
                "payment_due_date",
                "due_date",
                "_type",
                "title",
                "status",
            }:
                continue
            if key == "raw_text":
                details["raw_text_snippet"] = str(value)[:500] if value else None
                continue
            details[key] = value

        details["document_type"] = document_type
        details["_notes"] = notes

        output.append(
            {
                "title": title,
                "status": status,
                "details": details,
                "due_date": due_date,
            }
        )

    return output
